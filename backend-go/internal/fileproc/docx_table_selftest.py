#!/usr/bin/env python3
"""docx_table_selftest.py — normalize_tables 回归自检（基于真实 DOCX 产物形态）。

守护 2026-09-18 修复的「大量块不匹配」缺陷。用**真实 pdf2docx 产物的关键形态**做夹具：
  · tblGrid 是等宽占位符（2 列表恒为 50/50）；
  · tcW 才是从原 PDF 量出的真实列宽（26.9% / 73.1%）；
  · tblLayout=autofit、tblW=0/auto、行高 hRule=exact。
断言输出产物必须满足：
  1) tblLayout=fixed（不是 autofit）；
  2) **tblGrid 已按 tcW 真值重建** —— 列占比≈26.9/73.1，绝不是 50/50（旧实现会把
     「时间」列凭空放宽一倍，这是本次回归断言的核心）；
  3) **tcW 保留且与新栅格一致**（跨列单元格为被跨列宽之和）——tcW 是唯一几何真值，
     旧实现「删光 tcW + autofit」正是块不匹配的根因；
  4) 所有行高规则统一 atLeast（可撑开治裁剪，不写死高度）；
  5) 采不到 tcW 的表退回 autofit 兜底（修复前行为，不误伤非 pdf2docx 来源的表）。

用法：python3 docx_table_selftest.py   （需 python-docx；无则退出码 2 = 跳过）
由 internal/fileproc/docx_table_layout_test.go 经 go test 驱动。
"""
import os
import shutil
import sys
import tempfile

# 脚本与被测模块同目录，直接把脚本所在目录塞进 sys.path：go test 里是以 `python3
# docx_table_selftest.py` 相对路径启动的（工作目录=包目录），不依赖 PYTHONPATH 配置。
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    from docx import Document
    from docx.shared import Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except Exception as e:              # 无 python-docx：跳过（退出码 2），不算失败
    sys.stderr.write("SKIP: python-docx 不可用: %s\n" % e)
    sys.exit(2)

import docx_translate as dt

# 真实 pdf2docx 形态（实测 cache.docx 表 0）：
#   gridCol = [5417, 5417]（等宽占位 50/50）  tcW = [2828, 7696]（真实 26.9%/73.1%）
GRID_PLACEHOLDER = [5417, 5417]
TCW_REAL = [2828, 7696]

fails = []


# 断言器：失败只往模块级 fails 记账、不抛异常，于是一次运行能报出全部不符项（而不是
# 第一个 assert 就退出）。退出码统一由 main() 结尾按 fails 是否为空给出。
def ck(cond, msg):
    if not cond:
        fails.append(msg)


def build(path, grid_cols, tcw_cols, with_tcw=True, rows=3):
    """构造模拟 pdf2docx 产物的 DOCX：等宽 tblGrid 占位 + 真实 tcW + 固定行高 + autofit。"""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(3.17)
    tbl = doc.add_table(rows=rows, cols=len(grid_cols))
    grid = tbl._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), grid_cols):
        gc.set(qn('w:w'), str(w))
    for ri, row in enumerate(tbl.rows):
        for ci, c in enumerate(row.cells):
            c.text = "Cell %d-%d" % (ri, ci)
            tcPr = c._tc.get_or_add_tcPr()
            if with_tcw:
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:w'), str(tcw_cols[ci]))
                tcW.set(qn('w:type'), 'dxa')
        trPr = row._tr.get_or_add_trPr()
        h = OxmlElement('w:trHeight')
        h.set(qn('w:val'), '400')
        h.set(qn('w:hRule'), 'exact')                              # 写死高度 → 会裁剪
        trPr.append(h)
    tblPr = tbl._tbl.find(qn('w:tblPr'))
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'autofit')                               # 交付版本被改成 autofit
    tblPr.append(lay)
    doc.save(path)


def read(path):
    """读出产物关键 XML 事实（独立解析，不复用被测代码的读取逻辑）。"""
    d = Document(path)
    tbl = d.element.body.find(qn('w:tbl'))
    tblPr = tbl.find(qn('w:tblPr'))
    lay = tblPr.find(qn('w:tblLayout'))
    tw = tblPr.find(qn('w:tblW'))
    grid = tbl.find(qn('w:tblGrid'))
    cols = [int(gc.get(qn('w:w')) or 0) for gc in grid.findall(qn('w:gridCol'))]
    tcw = []
    for tc in tbl.find(qn('w:tr')).findall(qn('w:tc')):
        w = tc.find(qn('w:tcPr')).find(qn('w:tcW'))
        tcw.append(int(w.get(qn('w:w')) or 0) if w is not None else 0)
    rules = set()
    for trPr in tbl.iter(qn('w:trPr')):
        for h in trPr.findall(qn('w:trHeight')):
            rules.add(h.get(qn('w:hRule')))
    return {'layout': lay.get(qn('w:type')) if lay is not None else None,
            'tblW': (tw.get(qn('w:w')), tw.get(qn('w:type'))) if tw is not None else (None, None),
            'cols': cols, 'tcw': tcw, 'rules': rules}


def main():
    tmp = tempfile.mkdtemp(prefix="tblfix_")
    try:
        # —— 用例 1：pdf2docx 真实形态（等宽占位栅格 + 真实 tcW）——
        p = os.path.join(tmp, "real.docx")
        build(p, GRID_PLACEHOLDER, TCW_REAL)
        dt.normalize_tables(p)
        r = read(p)
        ck(r['layout'] == 'fixed', "layout=%s，应为 fixed（旧实现是 autofit）" % r['layout'])
        tot = float(sum(r['cols'])) or 1.0
        ratios = [c / tot for c in r['cols']]
        # ★ 核心回归断言：栅格必须按 tcW 真值重建（≈26.9/73.1），绝不是等宽占位 50/50
        ck(abs(ratios[0] - 0.269) < 0.01 and abs(ratios[1] - 0.731) < 0.01,
           "栅格列占比=%s，应≈[0.269, 0.731]（若≈[0.5,0.5] 说明误信了等宽占位栅格）" % ratios)
        ck(sum(r['cols']) == sum(TCW_REAL),
           "栅格总宽=%d，应为 tcW 总宽 %d（不得被拉伸到版心宽）" % (sum(r['cols']), sum(TCW_REAL)))
        # tcW 必须保留（旧实现删光）且与新栅格一致
        ck(len(r['tcw']) == 2 and sum(r['tcw']) > 0,
           "tcW 被删光（旧实现行为）：首行 tcW=%s" % r['tcw'])
        ck(r['tcw'] == [sum(r['cols'][:1]), sum(r['cols'][1:])],
           "tcW=%s 与新栅格 %s 不一致（两处口径必须统一）" % (r['tcw'], r['cols']))
        ck(r['rules'] == {'atLeast'},
           "行高规则应为 atLeast（可撑开、治裁剪），实际 %s" % (r['rules'] or "无"))
        ck(r['tblW'] == (str(sum(TCW_REAL)), 'dxa'),
           "tblW 应为 dxa=ΣtcW，实际 %s" % (r['tblW'],))

        # —— 用例 2：三列表真实几何 8.8% / 47.7% / 43.5%（源 PDF「3. 现状与痛点」表）——
        p2 = os.path.join(tmp, "three.docx")
        build(p2, [3611, 3611, 3611], [924, 5024, 4576])
        dt.normalize_tables(p2)
        r2 = read(p2)
        t2 = float(sum(r2['cols'])) or 1.0
        rr = [c / t2 for c in r2['cols']]
        ck(abs(rr[0] - 0.088) < 0.01 and abs(rr[1] - 0.477) < 0.01 and abs(rr[2] - 0.435) < 0.01,
           "三列栅格占比=%s，应≈[0.088,0.477,0.435]" % rr)
        ck(r2['layout'] == 'fixed', "三列表 layout=%s" % r2['layout'])

        # —— 用例 3：跨列单元格（gridSpan=2）—— tcW 应等于被跨列宽之和
        p3 = os.path.join(tmp, "span.docx")
        build(p3, GRID_PLACEHOLDER, TCW_REAL)
        d3 = Document(p3)
        tbl3 = d3.element.body.find(qn('w:tbl'))
        tc0 = tbl3.find(qn('w:tr')).findall(qn('w:tc'))[0]
        sp = OxmlElement('w:gridSpan')
        sp.set(qn('w:val'), '2')
        tc0.find(qn('w:tcPr')).append(sp)
        d3.save(p3)
        dt.normalize_tables(p3)
        r3 = read(p3)
        ck(r3['tcw'][0] == sum(r3['cols']),
           "跨列单元格 tcW=%d 应等于被跨列宽之和 %d" % (r3['tcw'][0], sum(r3['cols'])))

        # —— 用例 4：完全采不到 tcW → 退回 autofit 兜底 ——
        # 注意：python-docx 的 add_table 会自带 tcW（版心/列数），必须显式删光才是真「无 tcW」。
        p4 = os.path.join(tmp, "notcw.docx")
        build(p4, GRID_PLACEHOLDER, TCW_REAL)
        d4 = Document(p4)
        for el in list(d4.element.body.iter(qn('w:tcW'))):
            el.getparent().remove(el)
        d4.save(p4)
        dt.normalize_tables(p4)
        r4 = read(p4)
        ck(r4['layout'] == 'autofit', "无 tcW 表应退回 autofit，实际 %s" % r4['layout'])
        ck(r4['tblW'] == ('5000', 'pct'), "无 tcW 表 tblW 应为 5000/pct，实际 %s" % (r4['tblW'],))

        # —— 用例 5：水印切片清除（灰底上的「白色斑块」根因）——
        # 真实现场：pdf2docx 把原 PDF 极淡水印切成小图（全像素最暗通道 ≥214）嵌进单元格；
        # 真图（截图/图标）必有深色像素。判据：最深像素 <190 保留、≥190 移除。
        try:
            from PIL import Image as _Img
        except Exception:
            _Img = None
        if _Img is not None:
            p5 = os.path.join(tmp, "wm.docx")
            build(p5, GRID_PLACEHOLDER, TCW_REAL, rows=2)
            d5 = Document(p5)
            cell = d5.element.body.find(qn('w:tbl')).find(qn('w:tr')).findall(qn('w:tc'))[0]
            faint = _Img.new('RGB', (48, 28), (255, 255, 255))
            for x in range(0, 48, 4):                      # 极淡水印字形（灰 240）
                for y in range(4, 24, 6):
                    faint.putpixel((x, y), (240, 241, 243))
            real = _Img.new('RGB', (48, 28), (255, 255, 255))
            for x in range(10, 38):                        # 真图标（深色笔画）
                for y in range(8, 20):
                    real.putpixel((x, y), (30, 30, 30))
            pf, pr = os.path.join(tmp, "faint.png"), os.path.join(tmp, "real.png")
            faint.save(pf), real.save(pr)
            # 同一格嵌近白图+深色图（高层 API 走 part 关系解析）
            cell0 = d5.tables[0].rows[0].cells[0]
            cell0.paragraphs[-1].add_run().add_picture(pf, width=Cm(1.2))
            cell0.paragraphs[-1].add_run().add_picture(pr, width=Cm(1.2))
            d5.save(p5)
            n_before = len(list(d5.element.body.iter(qn('a:blip'))))
            removed = dt.strip_watermark_fragments(p5)
            d5b = Document(p5)
            n_after = len(list(d5b.element.body.iter(qn('a:blip'))))
            ck(n_before - n_after == 1,
               "水印切片清除数=%d（应恰为 1：近白图删、深色图留）" % removed)
            ck(n_after >= 1, "深色真图被误删：剩余 blip=%d" % n_after)

        # —— 用例 6：底色列内多数表决（水印误采的 dedfe3/ffffff 拼布 → 统一）——
        p6 = os.path.join(tmp, "shd.docx")
        build(p6, GRID_PLACEHOLDER, TCW_REAL, rows=3)
        d6 = Document(p6)
        tbl6 = d6.element.body.find(qn('w:tbl'))
        for ri, tc in enumerate(tbl6.findall(qn('w:tr'))):
            tc0 = tc.findall(qn('w:tc'))[0]
            tcPr = tc0.find(qn('w:tcPr'))
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'DEDFE3' if ri != 1 else 'FFFFFF')  # 灰、白、灰 拼布
            tcPr.append(shd)
        d6.save(p6)
        dt.unify_column_shading(p6)
        d6b = Document(p6)
        tbl6b = d6b.element.body.find(qn('w:tbl'))
        fills = []
        for tc in tbl6b.findall(qn('w:tr')):
            shd = tc.findall(qn('w:tc'))[0].find(qn('w:tcPr')).find(qn('w:shd'))
            fills.append(shd.get(qn('w:fill')) if shd is not None else None)
        ck(len(set(f.upper() for f in fills)) == 1,
           "拼布未统一：列内底色=%s（应为同一近白色）" % fills)
        ck(fills[0].upper() == 'DEDFE3',
           "表决结果应为多数色 DEDFE3，实际 %s" % fills[0])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if fails:
        for m in fails:
            sys.stderr.write("FAIL: %s\n" % m)
        sys.stderr.write("docx_table_selftest: %d 项断言失败\n" % len(fails))
        return 1
    print("docx_table_selftest OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
